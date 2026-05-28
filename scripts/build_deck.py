"""Build the Metromobile Cursor demo deck for UCLA MSBA students.

Goal: every slide explains a concept with a real example + has speaker notes.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(__file__).resolve().parent.parent / "Metromobile_Demo.pptx"

# --- Theme -----------------------------------------------------------------
INK = RGBColor(0x14, 0x1B, 0x2D)       # near-black for headings
BODY = RGBColor(0x33, 0x39, 0x47)      # softer body text
MUTED = RGBColor(0x6B, 0x72, 0x80)     # captions
ACCENT = RGBColor(0x2D, 0x6C, 0xDF)    # blue accents
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
CODE_BG = RGBColor(0xF3, 0xF4, 0xF6)
CODE_INK = RGBColor(0x1F, 0x29, 0x37)
DIVIDER = RGBColor(0xE5, 0xE7, 0xEB)

# --- Helpers ---------------------------------------------------------------

def add_slide(prs, title_text=None, subtitle=None, eyebrow=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Eyebrow
    if eyebrow:
        eb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.4))
        p = eb.text_frame.paragraphs[0]
        p.text = eyebrow.upper()
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT
    if title_text:
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(12.3), Inches(1.0))
        p = tb.text_frame.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = INK
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12.3), Inches(0.5))
        p = sb.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(16)
        p.font.color.rgb = MUTED
    return slide

def add_text(slide, left, top, width, height, runs, anchor=None):
    """runs = list of (text, size, bold, color, indent) tuples or a string."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    if isinstance(runs, str):
        p = tf.paragraphs[0]
        p.text = runs
        p.font.size = Pt(16)
        p.font.color.rgb = BODY
        return tb
    for i, item in enumerate(runs):
        if isinstance(item, str):
            text, size, bold, color, indent = item, 16, False, BODY, 0
        else:
            # Accept (text, size), (text, size, bold), (text, size, bold, color), or (text, size, bold, color, indent)
            text = item[0]
            size = item[1] if len(item) > 1 else 16
            bold = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else BODY
            indent = item[4] if len(item) > 4 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.level = int(indent)
        p.space_after = Pt(6)
    return tb

def bullets(slide, left, top, width, height, items, size=16, color=BODY, bullet_char="•"):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, indent = item
        else:
            text, indent = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        prefix = "    " * indent + (bullet_char + "  " if indent == 0 else "– ")
        p.text = prefix + text
        p.font.size = Pt(size - (2 if indent else 0))
        p.font.color.rgb = MUTED if indent else color
        p.space_after = Pt(6)
    return tb

def code_block(slide, left, top, width, height, code, size=12):
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    bg.fill.solid()
    bg.fill.fore_color.rgb = CODE_BG
    bg.line.fill.background()
    bg.adjustments[0] = 0.04
    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.1), Inches(width - 0.3), Inches(height - 0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line if line else " "
        p.font.name = "Consolas"
        p.font.size = Pt(size)
        p.font.color.rgb = CODE_INK
        p.space_after = Pt(0)
    return bg

def divider(slide, left, top, width):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = DIVIDER
    line.line.fill.background()

def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def callout(slide, left, top, width, height, label, body, color=ACCENT):
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xEF, 0xF4, 0xFE)
    bg.line.color.rgb = color
    bg.line.width = Pt(1)
    bg.adjustments[0] = 0.05
    tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.1), Inches(width - 0.4), Inches(height - 0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(14)
    p2.font.color.rgb = BODY

# --- Build the deck --------------------------------------------------------

def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ============================================================
    # Slide 1 — Title
    # ============================================================
    s = add_slide(prs)
    add_text(s, 0.6, 2.6, 12, 1.2, [("Cursor for Analytics", 54, True, INK)])
    add_text(s, 0.6, 3.9, 12, 1.0, [("Walking a $2M business decision with an AI teammate", 22, False, MUTED)])
    add_text(s, 0.6, 5.0, 12, 0.5, [("UCLA MSBA · 40 minutes · live demo", 14, False, ACCENT)])
    add_notes(s,
        "Welcome. Today I'll show you how an AI editor — Cursor — handles a real business question, "
        "and the few features (Rules, Skills, Commands, Plan mode) that turn it from a chatbot into "
        "a real analytics teammate. We're using a fake broadband company called Metromobile and 5 CSV files. "
        "By the end you'll have a checklist for using this in your own analytics work.")

    # ============================================================
    # Slide 2 — The scenario
    # ============================================================
    s = add_slide(prs, eyebrow="The scenario", title_text="Growth wants $2M for acquisition campaigns. Right call?")
    callout(s, 0.5, 2.0, 12.3, 1.2, "YOUR JOB AS THE NEW ANALYST",
            "Give leadership a data-backed answer using 5 CSV files. By the end of this demo, "
            "we'll have one — and you'll know how to do this on your own data.")
    add_text(s, 0.5, 3.5, 12, 0.5, [("Metromobile (synthetic broadband ISP) — what we have to work with:", 16, True, INK)])
    bullets(s, 0.7, 4.1, 12, 2.7, [
        "demo_customers.csv  →  50,000 subscribers, monthly revenue, customer lifetime value (CLV)",
        "demo_churns.csv  →  who left, when, and why (voluntary, non-payment, etc.)",
        "demo_calls.csv  →  200,000 contact-center calls with repeat-within-7-days flag",
        "demo_campaigns.csv  →  test vs control acquisition campaigns by ZIP and channel",
        "demo_market.csv  →  ZIP-level competitive context (does a fiber competitor exist?)",
    ], size=15)
    add_notes(s,
        "This is the question we'll answer with Cursor by the end. Same 5 CSV files, same question — "
        "but the answer will change as we layer Rules and Skills on top. "
        "Note: data is synthetic, ~50k subscribers, 6 months of events from Oct 2025 to Mar 2026.")

    # ============================================================
    # Slide 3 — What is Cursor
    # ============================================================
    s = add_slide(prs, eyebrow="The tool", title_text="What is Cursor?")
    add_text(s, 0.5, 1.9, 12.3, 0.7,
        [("An AI-powered code/analytics editor with an agent that reads your files, writes Python, runs it, and iterates.",
          18, False, BODY)])
    divider(s, 0.5, 3.0, 12.3)
    # 3 columns
    cols = [
        ("It IS", GREEN, [
            "A real editor (VS Code under the hood)",
            "An agent that can touch your files",
            "A loop: read → write code → run → fix → iterate",
        ]),
        ("It is NOT", RED, [
            "Just a chat window",
            "A magic black box (you tell it what to do)",
            "A replacement for understanding your data",
        ]),
        ("Today we use", ACCENT, [
            "Agent mode (does work for you)",
            "Rules (project context)",
            "Skills (analysis playbooks)",
            "Commands + Plan mode (bonus)",
        ]),
    ]
    for i, (header, c, items) in enumerate(cols):
        x = 0.5 + i * 4.27
        add_text(s, x, 3.2, 4.1, 0.5, [(header, 16, True, c)])
        bullets(s, x, 3.7, 4.1, 3.5, items, size=14)
    add_notes(s,
        "Think of Cursor as VS Code with an agent built in. The agent can read your CSVs, write a Python script, "
        "run it, look at the output, and fix it if something goes wrong. That last part — the loop — is what makes "
        "it 'agentic' rather than just chat. We're going to use Agent mode mostly. Ask mode is read-only Q&A, "
        "Plan mode designs work before doing it.")

    # ============================================================
    # Slide 4 — Three modes
    # ============================================================
    s = add_slide(prs, eyebrow="UI primer", title_text="Three chat modes you'll see today")
    rows = [
        ("Ask", "Read-only Q&A about your code or data", "“What columns does demo_calls.csv have?”", MUTED),
        ("Agent", "Does work: reads files, writes code, runs it, edits", "“Compute monthly churn rate and explain it.”", ACCENT),
        ("Plan", "Designs a multi-step plan, asks before executing", "“Plan an analysis to score subscribers by churn risk.”", GREEN),
    ]
    # Headers
    y0 = 2.1
    add_text(s, 0.5, y0, 1.8, 0.4, [("Mode", 14, True, MUTED)])
    add_text(s, 2.4, y0, 4.5, 0.4, [("What it does", 14, True, MUTED)])
    add_text(s, 7.0, y0, 5.8, 0.4, [("Example prompt", 14, True, MUTED)])
    divider(s, 0.5, y0 + 0.45, 12.3)
    for i, (mode, desc, ex, color) in enumerate(rows):
        y = y0 + 0.7 + i * 1.3
        add_text(s, 0.5, y, 1.8, 0.5, [(mode, 22, True, color)])
        add_text(s, 2.4, y, 4.5, 1.0, [(desc, 15, False, BODY)])
        code_block(s, 7.0, y - 0.05, 5.8, 0.85, ex, size=12)
    add_notes(s,
        "Most of our day is Agent. I'll switch to Plan once near the end to show how it changes the interaction. "
        "Ask is useful when you literally just want to look something up without the agent touching files. "
        "You change modes from the dropdown at the bottom of the chat panel.")

    # ============================================================
    # Slide 5 — 40 min plan
    # ============================================================
    s = add_slide(prs, eyebrow="Agenda", title_text="The next 40 minutes")
    rows = [
        ("0–5", "Setup", "Open the 5 CSVs. Quick tour of Cursor."),
        ("5–13", "Round 1: Raw Agent", "Ask the agent about churn with no extra context. Critique the answer."),
        ("13–20", "Round 2: Build a Rule live", "Add 5 lines of project context. Same prompt, sharper answer."),
        ("20–28", "Round 3: Build a Skill live", "Add a campaign-ROI playbook. Ask the $2M question."),
        ("28–33", "Round 4: Commands + Plan mode", "Make a /recap button. Use Plan mode on a complex task."),
        ("33–38", "Round 5: The $2M finale", "Load 3 more skills. Get the recommendation."),
        ("38–40", "Take-home", "Five things to do in your own projects this week."),
    ]
    y0 = 1.9
    add_text(s, 0.5, y0, 1.4, 0.4, [("Time", 13, True, MUTED)])
    add_text(s, 2.0, y0, 3.5, 0.4, [("Round", 13, True, MUTED)])
    add_text(s, 5.7, y0, 7.0, 0.4, [("What students see", 13, True, MUTED)])
    divider(s, 0.5, y0 + 0.4, 12.3)
    for i, (t, r, w) in enumerate(rows):
        y = y0 + 0.55 + i * 0.65
        add_text(s, 0.5, y, 1.4, 0.5, [(t, 14, True, ACCENT)])
        add_text(s, 2.0, y, 3.5, 0.5, [(r, 14, True, INK)])
        add_text(s, 5.7, y, 7.0, 0.5, [(w, 13, False, BODY)])
    add_notes(s,
        "Each round adds one Cursor capability and shows the answer changing. Same data, same question — "
        "different output as we layer in context.")

    # ============================================================
    # Slide 6 — Round 1: raw prompt
    # ============================================================
    s = add_slide(prs, eyebrow="Round 1 of 5 · Raw Agent", title_text="Just ask the AI — no rules, no skills")
    add_text(s, 0.5, 1.9, 12.3, 0.4, [("The prompt I'm about to paste:", 15, True, INK)])
    code_block(s, 0.5, 2.4, 12.3, 1.6,
        "How many customers churned in the last 3 months?\n"
        "What is the monthly churn rate? What is the total annual\n"
        "revenue at risk if this continues?", size=15)
    add_text(s, 0.5, 4.3, 12.3, 0.5, [("What I expect the agent to do:", 15, True, INK)])
    bullets(s, 0.7, 4.8, 12, 1.8, [
        "Load demo_churns.csv. Count churn events in the last 3 months. Maybe ~5,000.",
        "Divide by total customers to get a monthly rate (~3%).",
        "Multiply monthly revenue lost × 12 to get an 'annual revenue at risk' number.",
        "Probably treat all disconnect_types the same. Probably no caveats. Big scary headline.",
    ], size=14)
    add_notes(s,
        "I'm about to run this. Watch for: does it split voluntary vs non-payment? Does it caveat the annualization? "
        "Spoiler — without a Rule it usually doesn't. That's our motivation for Rules on the next slide.")

    # ============================================================
    # Slide 7 — What the raw answer misses
    # ============================================================
    s = add_slide(prs, eyebrow="Round 1 debrief", title_text="What the raw answer missed")
    # Two columns: what we got vs what's wrong with it
    add_text(s, 0.5, 1.9, 6.0, 0.4, [("Likely raw answer", 15, True, MUTED)])
    code_block(s, 0.5, 2.4, 6.0, 3.5,
        "~5,000 customers churned\n"
        "~3% monthly churn rate\n"
        "~$5M annual revenue at risk\n"
        "(if this continues)\n\n"
        "Suggestion: invest in\n"
        "retention to reduce churn.", size=13)
    add_text(s, 6.8, 1.9, 6.0, 0.4, [("What's wrong with it", 15, True, RED)])
    bullets(s, 6.8, 2.4, 6.0, 4.0, [
        "Lumps voluntary churn with non-payment writeoffs.",
        "Retention offers can't fix non-payment — that's a billing problem.",
        "The actionable slice is roughly half of that $5M.",
        "No mention of fiber competitors in the data.",
        "No annualization caveat (data is only 6 months).",
        "No assumptions, confidence, or risks stated.",
    ], size=14, color=BODY)
    add_notes(s,
        "Take a moment to read both columns. Ask the room: who would feel comfortable taking $5M of "
        "revenue-at-risk to their CFO? Nobody. The agent doesn't know how Metromobile thinks about churn "
        "yet. That's exactly what Rules are for — next slide.")

    # ============================================================
    # Slide 8 — Rules definition
    # ============================================================
    s = add_slide(prs, eyebrow="Round 2 of 5 · The fix", title_text="Rules — persistent context for your AI")
    add_text(s, 0.5, 1.9, 12.3, 0.6,
        [("A Rule is a small markdown file that loads into every chat in this project. ",
          18, False, BODY)])
    add_text(s, 0.5, 2.5, 12.3, 0.4,
        [("Like an employee handbook: written once, applied to every task.", 16, False, MUTED)])
    divider(s, 0.5, 3.1, 12.3)
    add_text(s, 0.5, 3.3, 6.0, 0.5, [("Good things to put in a Rule", 15, True, GREEN)])
    bullets(s, 0.5, 3.85, 6.0, 3.0, [
        "Definitions: 'voluntary churn' means disconnect_type='voluntary'",
        "Conventions: always annualize 6-mo data with ×2",
        "Style: end recommendations with assumptions + confidence",
        "Column names: repeat-call column is is_repeat_7d",
    ], size=14)
    add_text(s, 6.8, 3.3, 6.0, 0.5, [("Don't put in a Rule", 15, True, RED)])
    bullets(s, 6.8, 3.85, 6.0, 3.0, [
        "Step-by-step formulas (use a Skill instead)",
        "Anything you only need for one task",
        "Whole style guides — keep it short, < 50 lines",
        "Secrets, credentials, or large data dumps",
    ], size=14)
    add_notes(s,
        "Think of Rules as the things you'd put on a 1-page onboarding doc for a new analyst joining your team. "
        "Definitions, naming conventions, how to present numbers. Not the actual SOPs for specific analyses — "
        "those are Skills.")

    # ============================================================
    # Slide 9 — Building a rule
    # ============================================================
    s = add_slide(prs, eyebrow="Round 2 · Live coding", title_text="Building a Rule in 60 seconds")
    add_text(s, 0.5, 1.85, 12.3, 0.4, [("File: .cursor/rules/metromobile.mdc", 14, True, ACCENT)])
    code_block(s, 0.5, 2.3, 12.3, 4.3,
        "---\n"
        "description: Metromobile analytics conventions\n"
        "alwaysApply: true\n"
        "---\n\n"
        "# When analyzing Metromobile data\n\n"
        "- Lead with disconnect_type = 'voluntary'. That's what retention can fix.\n"
        "- Treat non_payment as a collections issue, not a retention target.\n"
        "- Period is Oct 2025 – Mar 2026 (6 months). Annualize with ×2 and label it.\n"
        "- Repeat-call column is is_repeat_7d.\n"
        "- End $ recommendations with assumptions, confidence, and one key risk.\n",
        size=13)
    add_text(s, 0.5, 6.8, 12.3, 0.4,
        [("Save the file → Cmd/Ctrl+Shift+P → Reload Window → New Chat → rerun the same prompt.", 13, False, MUTED)])
    add_notes(s,
        "I'll type this live. The frontmatter (between the dashes) tells Cursor what kind of rule this is. "
        "alwaysApply: true means it loads into every chat. The body is what the AI reads. "
        "Notice it's plain English — no special syntax. The biggest mistake people make is writing too much. "
        "Five lines beats fifty.")

    # ============================================================
    # Slide 10 — Watch the rule fire
    # ============================================================
    s = add_slide(prs, eyebrow="Round 2 · The reveal", title_text="Same prompt, with the Rule loaded")
    add_text(s, 0.5, 1.9, 6.0, 0.4, [("What changes in the answer", 15, True, GREEN)])
    bullets(s, 0.5, 2.4, 6.0, 4.5, [
        "Leads with voluntary churn (~2,100 in last 3 mo).",
        "Separates non-payment as a non-retention problem.",
        "Annualized number is labeled with ×2 assumption.",
        "Ends with confidence and one key risk.",
        "Same data. Same prompt. Sharper output.",
    ], size=15)
    add_text(s, 6.8, 1.9, 6.0, 0.4, [("Why this matters in real work", 15, True, ACCENT)])
    bullets(s, 6.8, 2.4, 6.0, 4.5, [
        "You stop repeating yourself in every chat.",
        "New teammates inherit your project's conventions automatically.",
        "Consistency across analyses — no more drift.",
        "5 lines of Rule replaces a 20-line system prompt you'd otherwise paste every time.",
    ], size=15)
    add_notes(s,
        "This is the 'aha' moment of the demo. The agent didn't get smarter — we just told it once how we think. "
        "Five lines, applied to every chat from now on.")

    # ============================================================
    # Slide 11 — Skills definition
    # ============================================================
    s = add_slide(prs, eyebrow="Round 3 of 5 · Going further", title_text="Skills — playbooks for specific analyses")
    add_text(s, 0.5, 1.9, 12.3, 0.7,
        [("A Skill is a multi-step recipe that loads only when the question matches. ",
          18, False, BODY)])
    add_text(s, 0.5, 2.55, 12.3, 0.5,
        [("Like an SOP: 'When asked about campaign ROI, run these 8 steps and output this table.'",
          16, False, MUTED)])
    divider(s, 0.5, 3.2, 12.3)
    # Rule vs Skill comparison
    add_text(s, 0.5, 3.4, 6.0, 0.5, [("Rules", 16, True, ACCENT)])
    bullets(s, 0.5, 3.95, 6.0, 3.0, [
        "Always on, in every chat",
        "Judgment & conventions",
        "Short (≤ 50 lines)",
        "Loaded automatically",
        "Analogy: employee handbook",
    ], size=14)
    add_text(s, 6.8, 3.4, 6.0, 0.5, [("Skills", 16, True, GREEN)])
    bullets(s, 6.8, 3.95, 6.0, 3.0, [
        "On-demand, when topic matches description",
        "Step-by-step workflows",
        "Can be longer (steps, formulas, output specs)",
        "Loaded when prompt triggers them",
        "Analogy: standard operating procedure",
    ], size=14)
    add_notes(s,
        "Both are markdown files. The difference is when they fire. Rules: every chat, always. "
        "Skills: only when the user's question matches the skill's description. "
        "Rule of thumb — if you'd want it on every chat about this project, it's a Rule. "
        "If it's a specific analysis with multiple steps, it's a Skill.")

    # ============================================================
    # Slide 12 — Building a skill
    # ============================================================
    s = add_slide(prs, eyebrow="Round 3 · Live coding", title_text="Building a Skill in 90 seconds")
    add_text(s, 0.5, 1.85, 12.3, 0.4, [("File: .cursor/skills/campaign-roi/SKILL.md", 14, True, ACCENT)])
    code_block(s, 0.5, 2.3, 12.3, 4.5,
        "---\n"
        "name: campaign-roi\n"
        "description: Use when the user asks about acquisition campaign ROI,\n"
        "  $2M, test vs control, connect rate, or cost per connect.\n"
        "---\n\n"
        "# Campaign ROI playbook\n\n"
        "## Steps\n"
        "1. Load demo_campaigns.csv. Compute connect rate by channel × group.\n"
        "2. Incremental lift = test rate − control rate.\n"
        "3. Test cost: door_to_door × $80 + digital_ads × $15.\n"
        "4. Project $2M at the same channel mix.\n"
        "5. Year-1 revenue = incr. connects × avg monthly_revenue × 12.\n\n"
        "## Required output\n"
        "Table: channel | control rate | test rate | lift | cost per incr. connect.\n"
        "Then $2M projection block. End with one-line verdict.\n",
        size=12)
    add_text(s, 0.5, 6.9, 12.3, 0.4,
        [("The 'description' is the trigger. Write it like the question your user would ask.", 13, False, MUTED)])
    add_notes(s,
        "The description is the most important line — that's what the agent matches against. "
        "Write it as the question the user would actually type. "
        "The body can be as detailed as you want, but cover: steps, required output format, and any unit costs.")

    # ============================================================
    # Slide 13 — Ask the $2M question
    # ============================================================
    s = add_slide(prs, eyebrow="Round 3 · The reveal", title_text="Now ask the $2M question")
    code_block(s, 0.5, 1.9, 12.3, 1.4,
        "Growth wants $2M for acquisition campaigns.\n"
        "Using the campaign ROI playbook, tell me what $2M actually buys\n"
        "and whether the unit economics support it.",
        size=15)
    add_text(s, 0.5, 3.6, 12.3, 0.4, [("Expected output", 15, True, GREEN)])
    bullets(s, 0.5, 4.05, 12.3, 2.8, [
        "Table: control vs test connect rates per channel, lift, cost per incremental connect",
        "$2M → ~530 incremental connects → ~$496k Year-1 gross revenue",
        "Year-1 net AFTER the $2M spend → negative",
        "Lifetime CLV recovery ≈ breakeven",
        "Verdict: don't lead with $2M acquisition on these unit economics",
    ], size=15)
    add_notes(s,
        "Without the skill, the agent would write you a 5-paragraph essay. With the skill, you get the table, "
        "the math, and the verdict in the same format every time. That's the value of Skills for repeatable analyses.")

    # ============================================================
    # Slide 14 — Commands + Plan mode
    # ============================================================
    s = add_slide(prs, eyebrow="Round 4 of 5 · Bonus capabilities", title_text="Commands & Plan mode")
    # Two columns
    # Commands
    add_text(s, 0.5, 1.9, 6.0, 0.5, [("Commands — one-shot buttons", 18, True, ACCENT)])
    add_text(s, 0.5, 2.5, 6.0, 1.0,
        [("Type /name in chat to trigger a saved instruction. Use for repeatable micro-tasks: "
          "summarize, format, extract.",
          14, False, BODY)])
    code_block(s, 0.5, 3.6, 6.0, 2.5,
        "# /recap\n\n"
        "Output exactly:\n"
        "Question: …\n"
        "Answer: …\n"
        "Why we believe it: …\n"
        "What to verify: …",
        size=12)
    add_text(s, 0.5, 6.25, 6.0, 0.4, [("Lives in .cursor/commands/recap.md", 12, False, MUTED)])
    # Plan mode
    add_text(s, 6.8, 1.9, 6.0, 0.5, [("Plan mode — design before execution", 18, True, GREEN)])
    add_text(s, 6.8, 2.5, 6.0, 1.0,
        [("Agent proposes the steps first, asks before doing the work. Use for multi-step or "
          "high-stakes tasks.",
          14, False, BODY)])
    code_block(s, 6.8, 3.6, 6.0, 2.5,
        "[Plan mode prompt]\n\n"
        "Plan a multi-step analysis\n"
        "to identify subscribers most\n"
        "at risk of voluntary churn\n"
        "in the next 60 days.\n\n"
        "Don't execute yet.",
        size=12)
    add_text(s, 6.8, 6.25, 6.0, 0.4, [("Switch via the chat mode dropdown.", 12, False, MUTED)])
    add_notes(s,
        "Commands cost 30 seconds to create and save you typing the same instruction over and over. "
        "Plan mode costs you 30 seconds of reading and saves you 20 minutes of fixing the wrong work. "
        "Use Plan when scope is ambiguous or stakes are high — like building a churn-risk score for "
        "the first time.")

    # ============================================================
    # Slide 15 — The Finale
    # ============================================================
    s = add_slide(prs, eyebrow="Round 5 of 5 · The finale", title_text="Where should the $2M actually go?")
    add_text(s, 0.5, 1.85, 12.3, 0.5,
        [("Now load three more pre-built skills and ask the question Growth is really asking:", 15, False, BODY)])
    code_block(s, 0.5, 2.4, 12.3, 1.6,
        "We have $2M. Compare three options: acquisition campaigns,\n"
        "voluntary churn prevention, and repeat-call reduction.\n"
        "Annual value, confidence, time to impact, key risk — then recommend.",
        size=14)
    add_text(s, 0.5, 4.2, 12.3, 0.5, [("Expected output", 15, True, GREEN)])
    # 3-row table preview
    rows = [
        ("Acquisition", "~$496k Year-1 gross · NEGATIVE Year-1 net", RED),
        ("Churn prevention (voluntary, fiber ZIPs)", "~$2.6–3.0M net annualized", GREEN),
        ("Repeat-call reduction", "~$5.1M total ($160k direct + ~$4.9M after 50% causal)", GREEN),
    ]
    y0 = 4.7
    for i, (option, value, color) in enumerate(rows):
        y = y0 + i * 0.55
        add_text(s, 0.7, y, 4.5, 0.5, [(option, 14, True, INK)])
        add_text(s, 5.4, y, 7.5, 0.5, [(value, 13, False, color)])
    add_notes(s,
        "This is the moment of the demo. Same 5 CSVs as the beginning, but the agent now has Metromobile's "
        "playbook in front of it. The recommendation comes out structured, with all three options compared. "
        "Punchline: don't lead with $2M acquisition. Mix churn prevention and repeat reduction. "
        "That's the gap Rules and Skills closed.")

    # ============================================================
    # Slide 16 — Take-home
    # ============================================================
    s = add_slide(prs, eyebrow="Take-home", title_text="Five things you'll do this week in your own projects")
    items = [
        ("1.", "Open your project in Cursor. Spend 10 minutes in Agent mode exploring the data — same as we did with Metromobile."),
        ("2.", "The third time you give the same instruction, stop. Turn it into a Rule (`.cursor/rules/yourname.mdc`)."),
        ("3.", "When you find yourself running the same analysis on different data, turn it into a Skill (`.cursor/skills/<name>/SKILL.md`)."),
        ("4.", "For your most frequent one-liner asks (summarize, recap, format), make a Command (`/yourname`)."),
        ("5.", "For complex multi-step work, use Plan mode first. Approve the plan. Then execute."),
    ]
    y0 = 2.0
    for i, (num, body) in enumerate(items):
        y = y0 + i * 0.85
        add_text(s, 0.5, y, 0.8, 0.7, [(num, 26, True, ACCENT)])
        add_text(s, 1.3, y + 0.08, 11.5, 0.8, [(body, 16, False, BODY)])
    add_text(s, 0.5, 6.6, 12.3, 0.5,
        [("Repo: github.com/ephemeralme/metromobile-demo-data · branch demo · everything you saw is in there.",
          14, False, ACCENT)])
    add_notes(s,
        "Send students to the repo. Everything we did is there — rules, skills, commands, the script, this deck. "
        "Encourage them to fork it and try the demo on their own data.")

    prs.save(OUT)
    print(f"Saved: {OUT}")
    print(f"Slides: {len(prs.slides)}")

if __name__ == "__main__":
    build()
