"""Build Metromobile pre-demo PowerPoint (no McKinsey branding)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = Path(__file__).resolve().parent.parent / "Metromobile_Cursor_PreDemo.pptx"

# Colors
DARK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
GRAY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_title(shape, text, size=32):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = DARK


def add_bullets(text_frame, items, size=18, level0=True):
    text_frame.clear()
    for i, item in enumerate(items):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK if level0 else GRAY
        p.space_after = Pt(8)


def slide_title_only(prs, title, subtitle=None):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12), Inches(1.2))
    set_title(box, title, 40)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.6), Inches(3.4), Inches(12), Inches(0.8))
        tf = sub.text_frame
        tf.text = subtitle
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.color.rgb = GRAY
    return slide


def slide_title_content(prs, title, bullets):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.9))
    set_title(tbox, title, 28)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.2), Inches(5.5))
    add_bullets(body.text_frame, bullets)
    return slide


def slide_two_column(prs, title, left_title, left_items, right_title, right_items):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.9))
    set_title(tbox, title, 28)
    for col, (ct, items, x) in enumerate(
        [(left_title, left_items, 0.5), (right_title, right_items, 6.8)]
    ):
        ht = slide.shapes.add_textbox(Inches(x), Inches(1.2), Inches(6), Inches(0.5))
        ht.text_frame.text = ct
        ht.text_frame.paragraphs[0].font.bold = True
        ht.text_frame.paragraphs[0].font.size = Pt(20)
        ht.text_frame.paragraphs[0].font.color.rgb = ACCENT
        bx = slide.shapes.add_textbox(Inches(x), Inches(1.7), Inches(6), Inches(4.8))
        add_bullets(bx.text_frame, items, 16)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title_only(
        prs,
        "Metromobile × Cursor",
        "Same data. Same questions. Better decisions with context.",
    )

    slide_title_content(
        prs,
        "Today's business question",
        [
            "Metromobile leadership has $2M to invest.",
            "Growth proposes: spend it on acquisition campaigns (door-to-door + digital).",
            "Leadership needs a data-backed answer: Is that the right use of $2M?",
            "We have five CSV files — customers, churn, calls, campaigns, market.",
            "In the live demo, Cursor will help us get from raw data → recommendation.",
        ],
    )

    slide_title_content(
        prs,
        "What we're demoing (and what we're not)",
        [
            "Demoing: Cursor Agent + CSV analytics + Rules + Skills",
            "Not demoing today: GitHub, deployment, parallel agents, MCP/API setup",
            "Synthetic demo data (~50,000 customers in the files; narrative may say 500k)",
            "You will see three passes over the same prompts — only the context layer changes",
        ],
    )

    slide_title_content(
        prs,
        "Three ways to guide the AI",
        [
            "1. Raw — prompts + data files only (no project context)",
            "2. Rules — always-on business context (how we interpret churn, calls, ROI)",
            "3. Skills — playbooks for specific analyses (formulas, tables, outputs)",
            "Analogy: Raw = new hire with files | Rules = employee handbook | Skills = SOPs",
        ],
    )

    slide_two_column(
        prs,
        "Rules vs Skills",
        "Rules (handbook)",
        [
            "Apply to every chat in the project",
            "Short, stable guardrails",
            "Example: lead with voluntary churn",
            "Example: non-payment = collections, not retention",
            "Example: correlation ≠ causation for calls vs churn",
            "Stored in: .cursor/rules/*.mdc",
        ],
        "Skills (playbooks)",
        [
            "Apply when the task matches",
            "Step-by-step + required tables",
            "Example: campaign ROI / $2M acquisition math",
            "Example: retention offer NPV",
            "Example: repeat-reduction case ($5.1M)",
            "Example: investment comparison + recommendation",
            "Stored in: .cursor/skills/<name>/SKILL.md",
        ],
    )

    slide_title_content(
        prs,
        "Rules — what ours contain (Metromobile)",
        [
            "Dataset scale: report actual row counts (50k customers)",
            "Churn: voluntary first; then full disconnect mix",
            "disconnect_type: voluntary = retention | non_payment = collections",
            "Calls: use is_repeat_7d; annualize 6 months with ×2",
            "Presentations: gross vs net, assumptions, confidence, risks",
        ],
    )

    slide_title_content(
        prs,
        "Skills — what ours contain (Metromobile)",
        [
            "metromobile-campaign-roi — test/control, incremental connects, $2M projection",
            "metromobile-churn-retention — fiber ZIPs, $200 offer, 30% save rate",
            "metromobile-call-center — repeat rate by queue, $8/call annual cost",
            "metromobile-repeat-reduction-case — $160k direct + ~$4.9M indirect",
            "metromobile-investment-compare — three options + where to put $2M",
        ],
    )

    slide_title_content(
        prs,
        "Live demo flow (~40 minutes)",
        [
            "Act 1 — Raw (10 min): churn at risk + call segments → messy headline",
            "Bridge (3 min): what rules are",
            "Act 2 — Rules (10 min): same churn question + disconnect breakdown",
            "Bridge (3 min): what skills are",
            "Act 3 — Skills (10 min): repeat-reduction case + $2M comparison table",
            "Close (2 min): recommendation — don't lead with acquisition",
        ],
    )

    slide_title_content(
        prs,
        "What should change between acts?",
        [
            "Act 1 → Act 2: Smaller voluntary churn $; collections called out separately",
            "Act 2 → Act 3: Full ROI tables; explicit $2M recommendation",
            "Expected finale: Prefer churn prevention + repeat reduction over acquisition",
            "Direction matters more than exact dollars — agent may vary slightly",
        ],
    )

    slide_title_content(
        prs,
        "Facilitator checklist",
        [
            "Branch: demo on github.com/ephemeralme/metromobile-demo-data",
            "Before each act: .\\scripts\\demo-mode.ps1 raw | rules | skills",
            "Then: Reload Window + New Chat (critical)",
            "Follow prompts in DEMO_SCRIPT.md; keep ANSWER_KEY.md for yourself only",
            "Opening line: Growth wants $2M for acquisition — we'll decide if that's right",
        ],
    )

    slide_title_only(prs, "Let's run the demo", "Act 1: Raw mode")

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
